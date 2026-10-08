"""
questions.json → 確認用 Excel（難易度別シート）

実行: python scripts/export_review_excel.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from kidgame.data.models import Difficulty
from kidgame.data.review_csv import REVIEW_CSV_HEADERS, question_in_mode, question_to_review_row

ROOT = Path(__file__).resolve().parents[1]
JSON_PATH = ROOT / "data" / "questions.json"
OUT_PATH = ROOT / "data" / "questions_review.xlsx"

OPT_COL_START = 3
ANS_SYM_COL = 6
ANS_TEXT_COL = 7

SHEET_TITLE = {
    Difficulty.EASY: "イージー",
    Difficulty.NORMAL: "ノーマル",
    Difficulty.HARD: "ハード",
}


def _style_sheet(ws, questions: list[dict], rows: list[list]) -> None:
    header_fill = PatternFill("solid", fgColor="2F4F6F")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    correct_fill = PatternFill("solid", fgColor="C6EFCE")
    correct_opt_fill = PatternFill("solid", fgColor="E2F0D9")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(list(REVIEW_CSV_HEADERS))
    for col in range(1, len(REVIEW_CSV_HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for q_idx, q in enumerate(questions):
        row_data = rows[q_idx]
        ws.append(row_data)
        excel_row = q_idx + 2
        correct_col = OPT_COL_START

        for col in range(1, len(REVIEW_CSV_HEADERS) + 1):
            cell = ws.cell(row=excel_row, column=col)
            cell.alignment = wrap
            cell.border = border
            if col in (ANS_SYM_COL, ANS_TEXT_COL):
                cell.fill = correct_fill
                cell.font = Font(bold=True)
            if col == correct_col:
                cell.fill = correct_opt_fill
                cell.font = Font(bold=True)

    widths = {
        1: 14,
        2: 48,
        3: 20,
        4: 12,
        5: 12,
        6: 6,
        7: 20,
        8: 36,
        9: 52,
    }
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = (
            f"A1:{get_column_letter(len(REVIEW_CSV_HEADERS))}{len(rows) + 1}"
        )


def main() -> None:
    data = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    all_questions = data["questions"]

    wb = Workbook()
    wb.remove(wb.active)

    for mode in Difficulty:
        qs = [q for q in all_questions if question_in_mode(q, mode)]
        rows = [question_to_review_row(q) for q in qs]
        ws = wb.create_sheet(title=SHEET_TITLE[mode])
        _style_sheet(ws, qs, rows)
        print(f"Sheet {ws.title}: {len(qs)} rows")

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
